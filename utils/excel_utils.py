# ============================================================================
# EXCEL UTILITIES - JOB DATA STORAGE AND FILTERING
# ============================================================================
# Handles Excel file operations for job data:
# - Creating Excel workbook with headers
# - Loading existing jobs from Excel
# - Saving new jobs and avoiding duplicates
# - Filtering irrelevant jobs based on title keywords
# - Updating job application status after applying

from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from config import (
    EXCLUSION_TITLE_KEYWORDS,
    OUTPUT_FILE,
    OUTPUT_FOLDER,
    RELEVANT_TITLE_KEYWORDS,
)


class ExcelJobsRepository:
    """Manages job data in Excel file."""
    
    # Column headers for the Excel file
    HEADERS = [
        "Job Title",
        "Company",
        "Location",
        "Job URL",
        "Application Status",
        "Application Date",
    ]

    def __init__(self, output_file=OUTPUT_FILE):
        """
        Initialize with output file path.
        
        Args:
            output_file: Path to Excel file (default: from config)
        """
        self.output_file = output_file

    def create_if_required(self):
        """
        Create Excel file if it doesn't exist.
        Creates directory, workbook, adds headers, and sets column widths.
        """
        OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

        if self.output_file.exists():
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Naukri Jobs"
        sheet.append(self.HEADERS)

        # Make header row bold
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        self._set_widths(sheet)
        workbook.save(self.output_file)
        workbook.close()

    def load_jobs(self):
        """
        Load all jobs from Excel file.
        
        Returns:
            list: List of job dictionaries with all columns
        """
        self.create_if_required()

        workbook = load_workbook(self.output_file)
        sheet = workbook.active
        headers = self._headers(sheet)

        jobs = []
        for row in range(2, sheet.max_row + 1):
            url = sheet.cell(row=row, column=headers["Job URL"]).value

            # Skip empty rows
            if not url:
                continue

            jobs.append(
                {
                    "row": row,
                    "title": self._cell_text(sheet, row, headers["Job Title"]),
                    "company": self._cell_text(sheet, row, headers["Company"]),
                    "location": self._cell_text(sheet, row, headers["Location"]),
                    "url": str(url).strip(),
                    "status": self._cell_text(
                        sheet, row, headers["Application Status"]
                    ),
                    "date": self._cell_text(sheet, row, headers["Application Date"]),
                }
            )

        workbook.close()
        return jobs

    def save_new_jobs(self, jobs):
        """
        Add new jobs to Excel file, avoiding duplicates by URL.
        
        Args:
            jobs: List of job dictionaries to save
            
        Returns:
            int: Number of new jobs added
        """
        self.create_if_required()

        workbook = load_workbook(self.output_file)
        sheet = workbook.active
        headers = self._headers(sheet)
        url_column = headers["Job URL"]
        existing_urls = set()

        # Collect existing URLs
        for row in range(2, sheet.max_row + 1):
            value = sheet.cell(row=row, column=url_column).value
            if value:
                existing_urls.add(str(value).strip())

        added = 0
        for job in jobs:
            url = job["url"]

            # Skip if URL already exists (prevent duplicates)
            if url in existing_urls:
                continue

            sheet.append(
                [
                    job["title"],
                    job["company"],
                    job["location"],
                    url,
                    job.get("status", "Not Processed"),
                    job.get("date", ""),
                ]
            )
            existing_urls.add(url)
            added += 1

        self._set_widths(sheet)
        workbook.save(self.output_file)
        workbook.close()
        return added

    def filter_irrelevant_jobs(self):
        """
        Remove jobs from Excel that don't match relevance criteria.
        Keeps only jobs that:
        1. Contain at least one RELEVANT_TITLE_KEYWORD
        2. Don't contain any EXCLUSION_TITLE_KEYWORDS
        
        Returns:
            tuple: (kept_count, deleted_count) - number of jobs kept and deleted
        """
        self.create_if_required()

        workbook = load_workbook(self.output_file)
        sheet = workbook.active
        headers = self._headers(sheet)
        title_column = headers["Job Title"]

        deleted_count = 0
        kept_count = 0

        # Process from bottom to top to avoid row index shifts when deleting
        for row in range(sheet.max_row, 1, -1):
            title = self._cell_text(sheet, row, title_column)
            title_lower = title.lower()

            # Check if title contains relevant keywords
            is_relevant = any(
                keyword in title_lower for keyword in RELEVANT_TITLE_KEYWORDS
            )
            # Check if title contains exclusion keywords
            is_excluded = any(
                keyword in title_lower for keyword in EXCLUSION_TITLE_KEYWORDS
            )

            # Keep if relevant and not excluded
            if title and is_relevant and not is_excluded:
                kept_count += 1
            else:
                sheet.delete_rows(row)
                deleted_count += 1

        workbook.save(self.output_file)
        workbook.close()
        return kept_count, deleted_count

    def update_status(self, job, status):
        """
        Update application status for a job and record the timestamp.
        
        Args:
            job: Job dictionary (must contain 'url' key)
            status: Application result status to record
        """
        workbook = load_workbook(self.output_file)
        sheet = workbook.active
        headers = self._headers(sheet)
        target_url = job["url"]
        target_row = None

        # Find row with matching URL
        for row in range(2, sheet.max_row + 1):
            value = sheet.cell(row=row, column=headers["Job URL"]).value

            if value and str(value).strip() == target_url:
                target_row = row
                break

        if target_row is not None:
            # Update status and timestamp
            sheet.cell(row=target_row, column=headers["Application Status"]).value = (
                status
            )
            sheet.cell(row=target_row, column=headers["Application Date"]).value = (
                datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            )

        workbook.save(self.output_file)
        workbook.close()

    def _headers(self, sheet):
        """
        Get header column mapping from first row.
        
        Args:
            sheet: Openpyxl sheet object
            
        Returns:
            dict: Mapping of header name to column number
            
        Raises:
            ValueError: If required headers are missing
        """
        headers = {}
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row=1, column=column).value
            if value:
                headers[str(value).strip()] = column

        # Validate all required headers exist
        missing = [header for header in self.HEADERS if header not in headers]
        if missing:
            raise ValueError(f"Missing Excel headers: {', '.join(missing)}")

        return headers

    def _set_widths(self, sheet):
        """
        Set column widths for better readability.
        
        Args:
            sheet: Openpyxl sheet object
        """
        widths = {
            "A": 45,  # Job Title
            "B": 35,  # Company
            "C": 25,  # Location
            "D": 80,  # Job URL
            "E": 40,  # Application Status
            "F": 25,  # Application Date
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

    def _cell_text(self, sheet, row, column):
        """
        Get cell text value, safely handling None values.
        
        Args:
            sheet: Openpyxl sheet object
            row: Row number
            column: Column number
            
        Returns:
            str: Cell text or empty string
        """
        return str(sheet.cell(row=row, column=column).value or "").strip()
